"""MCP tool implementations for intake exploration, secrets, and deliverable Excel."""
from __future__ import annotations

import json
import logging
import re
import shutil
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from src.intake.extraction.document_extractor import DocumentExtractor, SENSITIVE_FIELD_HINTS
from src.intake.intake_index import IntakeIndex, build_intake_index
from src.services.run_secret_store import RunSecretStore

logger = logging.getLogger("aptp_intake_tools")

_REPO_ROOT = Path(__file__).resolve().parents[2]
TEMPLATE_PATH = _REPO_ROOT / "src" / "intake" / "templates" / "deliverables" / "set-de-pruebas-template.xlsx"
ANCHOR_MAP_PATH = _REPO_ROOT / "src" / "intake" / "templates" / "deliverables" / "anchor_map.json"


@dataclass
class IntakeToolContext:
    project_id: str
    run_id: str
    documents: List[Dict[str, Any]] = field(default_factory=list)
    index: Optional[IntakeIndex] = None
    artifact_dir: Optional[Path] = None
    asset_bindings: List[Dict[str, Any]] = field(default_factory=list)
    audit_log: List[Dict[str, Any]] = field(default_factory=list)

    def __post_init__(self) -> None:
        if self.index is None and self.documents:
            self.index = build_intake_index(self.project_id, self.documents)

    @property
    def secret_store(self) -> RunSecretStore:
        return RunSecretStore.for_run(self.run_id)


class IntakeTools:
    """In-process MCP tool surface for agents."""

    TOOL_NAMES = [
        "list_intake_files",
        "describe_document",
        "list_excel_sheets",
        "read_excel_range",
        "read_pdf_pages",
        "search_intake_text",
        "find_file_by_hint",
        "identify_sensitive_fields",
        "register_runtime_secret",
        "get_runtime_secret",
        "propose_test_asset_binding",
        "list_qa_obligations",
        "build_scenario_outline",
        "get_deliverable_template_info",
        "clone_deliverable_template",
        "write_deliverable_cells",
        "append_deliverable_test_rows",
        "validate_deliverable_structure",
    ]

    def __init__(self, ctx: IntakeToolContext):
        self.ctx = ctx
        self._extractor = DocumentExtractor()
        self._doc_by_id: Dict[str, Dict[str, Any]] = {}
        for doc in ctx.documents:
            fid = doc.get("id") or doc.get("drive_file_id") or doc.get("filename", "")
            self._doc_by_id[fid] = doc
            self._doc_by_id[doc.get("filename", "")] = doc

    def call(self, tool_name: str, arguments: Dict[str, Any]) -> Dict[str, Any]:
        handler = getattr(self, f"_tool_{tool_name}", None)
        if not handler:
            return {"error": f"Unknown tool: {tool_name}"}
        try:
            result = handler(**arguments)
            self.ctx.audit_log.append({"tool": tool_name, "args_keys": list(arguments.keys()), "ok": True})
            return result
        except Exception as exc:
            logger.exception("Tool %s failed", tool_name)
            self.ctx.audit_log.append({"tool": tool_name, "error": str(exc), "ok": False})
            return {"error": str(exc)}

    def _resolve_doc(self, file_id: str) -> Optional[Dict[str, Any]]:
        return self._doc_by_id.get(file_id)

    def _tool_list_intake_files(self) -> Dict[str, Any]:
        files = []
        for doc in self.ctx.documents:
            extra = doc.get("extra") or {}
            files.append(
                {
                    "file_id": doc.get("id") or doc.get("drive_file_id"),
                    "filename": doc.get("filename"),
                    "relative_path": doc.get("relative_path", ""),
                    "intake_role": extra.get("intake_role", "SUPPLEMENTARY"),
                    "format": extra.get("format"),
                    "size_hint": len(doc.get("raw_text") or ""),
                }
            )
        return {"files": files, "count": len(files)}

    def _tool_describe_document(self, file_id: str) -> Dict[str, Any]:
        doc = self._resolve_doc(file_id)
        if not doc:
            return {"error": f"File not found: {file_id}"}
        extra = doc.get("extra") or {}
        text = doc.get("raw_text") or ""
        return {
            "file_id": file_id,
            "filename": doc.get("filename"),
            "format": extra.get("format"),
            "intake_role": extra.get("intake_role"),
            "page_count": extra.get("metadata", {}).get("page_count"),
            "sheet_count": extra.get("metadata", {}).get("sheet_count"),
            "sample_headers": self._sample_headers(doc),
            "signals": {
                "has_passwords": extra.get("metadata", {}).get("has_password_hints", False),
                "has_steps": extra.get("metadata", {}).get("has_steps", False),
                "has_urls": extra.get("metadata", {}).get("has_urls", False),
            },
            "text_preview": text[:1500],
        }

    def _sample_headers(self, doc: Dict[str, Any]) -> List[str]:
        extra = doc.get("extra") or {}
        sheets = extra.get("sheets") or []
        headers = []
        for sheet in sheets[:3]:
            headers.extend(sheet.get("headers", [])[:10])
        return headers[:20]

    def _tool_list_excel_sheets(self, file_id: str) -> Dict[str, Any]:
        doc = self._resolve_doc(file_id)
        if not doc:
            return {"error": f"File not found: {file_id}"}
        extra = doc.get("extra") or {}
        sheets = extra.get("sheets") or []
        result = []
        for sheet in sheets:
            dims = {"rows": sheet.get("row_count", 0), "cols": sheet.get("col_count", 0)}
            preview = (sheet.get("rows") or [])[:3]
            result.append(
                {
                    "name": sheet.get("name"),
                    "rows": dims.get("rows", 0),
                    "cols": dims.get("cols", 0),
                    "preview": preview,
                }
            )
        return {"sheets": result}

    def _tool_read_excel_range(self, file_id: str, sheet: str, range_a1: str = "A1:Z100") -> Dict[str, Any]:
        doc = self._resolve_doc(file_id)
        if not doc:
            return {"error": f"File not found: {file_id}"}
        extra = doc.get("extra") or {}
        sheets = {s.get("name"): s for s in (extra.get("sheets") or [])}
        sheet_data = sheets.get(sheet)
        if not sheet_data:
            return {"error": f"Sheet not found: {sheet}"}
        rows = sheet_data.get("rows") or []
        start_row, end_row, start_col, end_col = self._parse_range(range_a1, len(rows))
        sliced = []
        for ri in range(start_row - 1, min(end_row, len(rows))):
            row = rows[ri]
            sliced.append(row[start_col - 1 : end_col])
        return {"sheet": sheet, "range": range_a1, "rows": sliced}

    def _tool_read_pdf_pages(self, file_id: str, pages: List[int]) -> Dict[str, Any]:
        doc = self._resolve_doc(file_id)
        if not doc:
            return {"error": f"File not found: {file_id}"}
        extra = doc.get("extra") or {}
        page_data = {p.get("page_number", p.get("page")): p for p in (extra.get("pages") or [])}
        result = []
        for p in pages:
            entry = page_data.get(p)
            if entry:
                result.append({"page": p, "text": entry.get("text", "")[:8000]})
        if not result and doc.get("raw_text"):
            result.append({"page": 1, "text": doc["raw_text"][:8000]})
        return {"pages": result}

    def _tool_search_intake_text(self, query: str, limit: int = 20) -> Dict[str, Any]:
        if not self.ctx.index:
            return {"hits": []}
        return {"hits": self.ctx.index.search(query, limit=limit)}

    def _tool_find_file_by_hint(self, hint: str) -> Dict[str, Any]:
        hint_norm = hint.replace("\\", "/").lower()
        basename = Path(hint_norm).name
        matches = []
        for doc in self.ctx.documents:
            rel = (doc.get("relative_path") or doc.get("filename", "")).lower()
            fname = doc.get("filename", "").lower()
            if basename in rel or basename in fname or hint_norm in rel:
                matches.append(
                    {
                        "file_id": doc.get("id") or doc.get("drive_file_id"),
                        "filename": doc.get("filename"),
                        "relative_path": doc.get("relative_path"),
                    }
                )
        return {"matches": matches}

    def _tool_identify_sensitive_fields(self, file_id: str, sheet: Optional[str] = None) -> Dict[str, Any]:
        doc = self._resolve_doc(file_id)
        if not doc:
            return {"error": f"File not found: {file_id}"}
        extra = doc.get("extra") or {}
        fields = []
        for sh in extra.get("sheets") or []:
            if sheet and sh.get("name") != sheet:
                continue
            headers = sh.get("headers") or ((sh.get("rows") or [[]])[0] if sh.get("rows") else [])
            for ci, header in enumerate(headers):
                if SENSITIVE_FIELD_HINTS.search(str(header)):
                    fields.append(
                        {
                            "sheet": sh.get("name"),
                            "column": get_column_letter(ci + 1),
                            "header": header,
                            "sensitive": True,
                        }
                    )
        return {"fields": fields, "count": len(fields)}

    def _tool_register_runtime_secret(self, key: str, value: str, source: str = "") -> Dict[str, Any]:
        self.ctx.secret_store.register(key, value, source=source)
        return {"registered": True, "key": key}

    def _tool_get_runtime_secret(self, key: str) -> Dict[str, Any]:
        val = self.ctx.secret_store.get(key)
        if val is None:
            return {"found": False, "key": key}
        return {"found": True, "key": key, "value": val}

    def _tool_propose_test_asset_binding(self, obligation_id: str, asset_keys: List[str]) -> Dict[str, Any]:
        binding = {"obligation_id": obligation_id, "asset_keys": asset_keys}
        self.ctx.asset_bindings.append(binding)
        return {"binding": binding}

    def _tool_list_qa_obligations(self) -> Dict[str, Any]:
        from src.intake.qa_spec_chunker import chunk_documents, obligations_from_chunks

        chunks = chunk_documents(self.ctx.documents)
        rows = obligations_from_chunks(chunks)
        return {"obligations": rows, "count": len(rows), "chunk_count": len(chunks)}

    def _tool_build_scenario_outline(
        self,
        target_url: str = "",
        bank_portal: bool = False,
        include_manual: bool = True,
    ) -> Dict[str, Any]:
        from src.automation.scenario_coverage import build_scenarios_from_obligations
        from src.intake.qa_spec_chunker import chunk_documents, obligations_from_chunks

        rows = obligations_from_chunks(chunk_documents(self.ctx.documents))
        scenarios = build_scenarios_from_obligations(
            rows,
            target_url or "https://example.com",
            bank_portal=bank_portal,
            include_manual=include_manual,
        )
        return {
            "scenarios": scenarios,
            "count": len(scenarios),
            "obligation_count": len(rows),
            "coverage_complete": len(scenarios) >= len(rows),
        }

    def _tool_get_deliverable_template_info(self) -> Dict[str, Any]:
        anchor_map = self._load_anchor_map()
        info = {"template_path": str(TEMPLATE_PATH), "exists": TEMPLATE_PATH.is_file()}
        if TEMPLATE_PATH.is_file():
            wb = load_workbook(TEMPLATE_PATH, read_only=True)
            info["sheets"] = wb.sheetnames
            info["merge_counts"] = {}
            wb.close()
            wb = load_workbook(TEMPLATE_PATH)
            for name in wb.sheetnames:
                info["merge_counts"][name] = len(list(wb[name].merged_cells.ranges))
            wb.close()
        info["anchors"] = anchor_map.get("anchors", {})
        return info

    def _tool_clone_deliverable_template(self, dest_filename: str = "deliverable.xlsx") -> Dict[str, Any]:
        if not TEMPLATE_PATH.is_file():
            return {"error": "Template not found"}
        dest_dir = self.ctx.artifact_dir or Path("generated-test-projects") / self.ctx.project_id / "aptp-artifacts" / self.ctx.run_id
        dest_dir.mkdir(parents=True, exist_ok=True)
        dest = dest_dir / dest_filename
        shutil.copy2(TEMPLATE_PATH, dest)
        return {"path": str(dest), "filename": dest_filename}

    def _tool_write_deliverable_cells(self, path: str, cells: Dict[str, Any]) -> Dict[str, Any]:
        wb = load_workbook(path)
        written = 0
        for ref, value in cells.items():
            if "!" in ref:
                sheet_name, cell = ref.split("!", 1)
                if sheet_name in wb.sheetnames:
                    wb[sheet_name][cell] = value
                    written += 1
            else:
                for ws in wb.worksheets:
                    ws[ref] = value
                    written += 1
                    break
        wb.save(path)
        return {"written": written}

    def _tool_append_deliverable_test_rows(self, path: str, rows: List[Dict[str, Any]], start_row: int = 10) -> Dict[str, Any]:
        wb = load_workbook(path)
        sheet_name = "Casos de Pruebas"
        if sheet_name not in wb.sheetnames:
            sheet_name = wb.sheetnames[1] if len(wb.sheetnames) > 1 else wb.sheetnames[0]
        ws = wb[sheet_name]
        ri = start_row
        for row in rows:
            ws.cell(row=ri, column=3, value=row.get("hierarchy", row.get("title", "")))
            ws.cell(row=ri, column=4, value=row.get("steps", ""))
            ws.cell(row=ri, column=5, value=row.get("expected", ""))
            ws.cell(row=ri, column=6, value=row.get("result", ""))
            ri += 1
        wb.save(path)
        return {"appended": len(rows), "start_row": start_row}

    def _tool_validate_deliverable_structure(self, path: str) -> Dict[str, Any]:
        if not TEMPLATE_PATH.is_file():
            return {"valid": False, "error": "Template missing"}
        anchor_map = self._load_anchor_map()
        expected_sheets = anchor_map.get("sheet_names", ["Datos Generales", "Casos de Pruebas", "Aprobaciones"])
        wb = load_workbook(path)
        actual_sheets = wb.sheetnames
        sheet_ok = actual_sheets == expected_sheets
        merge_diff = {}
        tpl = load_workbook(TEMPLATE_PATH)
        for name in expected_sheets:
            if name in wb.sheetnames and name in tpl.sheetnames:
                merge_diff[name] = len(list(wb[name].merged_cells.ranges)) - len(list(tpl[name].merged_cells.ranges))
        wb.close()
        tpl.close()
        anchor_hits = {}
        for anchor in anchor_map.get("structural_anchors", []):
            sheet, cell, text = anchor["sheet"], anchor["cell"], anchor["text"]
            if sheet in load_workbook(path).sheetnames:
                val = str(load_workbook(path, read_only=True)[sheet][cell].value or "")
                anchor_hits[f"{sheet}!{cell}"] = text.lower() in val.lower()
        return {
            "valid": sheet_ok and all(v == 0 for v in merge_diff.values()),
            "sheets_match": sheet_ok,
            "merge_diff": merge_diff,
            "anchor_hits": anchor_hits,
        }

    def _load_anchor_map(self) -> Dict[str, Any]:
        if ANCHOR_MAP_PATH.is_file():
            return json.loads(ANCHOR_MAP_PATH.read_text(encoding="utf-8"))
        return {
            "sheet_names": ["Datos Generales", "Casos de Pruebas", "Aprobaciones"],
            "structural_anchors": [
                {"sheet": "Datos Generales", "cell": "A1", "text": "FORMATO SET DE PRUEBAS"},
                {"sheet": "Aprobaciones", "cell": "A1", "text": "APROBACIONES"},
            ],
            "anchors": {},
        }

    @staticmethod
    def _parse_range(range_a1: str, max_rows: int) -> tuple:
        m = re.match(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", range_a1.upper())
        if not m:
            return 1, min(100, max_rows), 1, 26
        from openpyxl.utils import column_index_from_string

        sc, sr, ec, er = m.group(1), int(m.group(2)), m.group(3), int(m.group(4))
        return sr, er, column_index_from_string(sc), column_index_from_string(ec)


def get_intake_tools(ctx: IntakeToolContext) -> IntakeTools:
    return IntakeTools(ctx)
