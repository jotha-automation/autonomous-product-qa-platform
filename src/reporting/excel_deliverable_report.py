"""Deterministic Excel deliverable writer — clones E59 template structure."""
from __future__ import annotations

import json
import re
import shutil
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple

from src.intake.test_naming import (
    assign_case_codes,
    automation_tier_label_es,
    format_case_code,
    format_test_set_name,
    infer_test_set_type,
)

_REPO_ROOT = Path(__file__).resolve().parents[1]
_TEMPLATE = _REPO_ROOT / "intake" / "templates" / "deliverables" / "set-de-pruebas-template.xlsx"
_ANCHOR_MAP = _REPO_ROOT / "intake" / "templates" / "deliverables" / "anchor_map.json"

# Playwright / terminal output may include ANSI escapes and control chars illegal in XLSX.
_ANSI_ESCAPE_RE = re.compile(r"\x1b\[[0-9;]*[A-Za-z]")
_ILLEGAL_XLSX_CHAR_RE = re.compile(r"[\000-\010]|[\013-\014]|[\016-\037]")


def sanitize_excel_text(value: Any, *, max_len: int = 32000) -> str:
    """Strip ANSI color codes and XML-illegal control chars before writing cells."""
    text = str(value or "")
    text = _ANSI_ESCAPE_RE.sub("", text)
    text = _ILLEGAL_XLSX_CHAR_RE.sub("", text)
    if len(text) > max_len:
        text = text[: max_len - 1] + "…"
    return text


def _set_cell_value(ws, cell_ref: str, value: Any) -> None:
    _writable_cell(ws, cell_ref).value = sanitize_excel_text(value) or None


def load_anchor_map() -> Dict[str, Any]:
    if _ANCHOR_MAP.is_file():
        return json.loads(_ANCHOR_MAP.read_text(encoding="utf-8"))
    return {}


def clone_template(dest: Path) -> Path:
    dest.parent.mkdir(parents=True, exist_ok=True)
    if not _TEMPLATE.is_file():
        raise FileNotFoundError(f"Deliverable template missing: {_TEMPLATE}")
    shutil.copy2(_TEMPLATE, dest)
    return dest


def _writable_cell(ws, cell_ref: str):
    """Return the top-left cell when *cell_ref* sits inside a merged range."""
    from openpyxl.cell.cell import MergedCell
    from openpyxl.utils import column_index_from_string
    from openpyxl.utils.cell import coordinate_from_string

    col_letter, row = coordinate_from_string(cell_ref)
    col = column_index_from_string(col_letter)
    cell = ws.cell(row=row, column=col)
    if isinstance(cell, MergedCell):
        for merged in ws.merged_cells.ranges:
            if merged.min_row <= row <= merged.max_row and merged.min_col <= col <= merged.max_col:
                return ws.cell(row=merged.min_row, column=merged.min_col)
    return cell


def _normalize_inc_number(inc: Any) -> str:
    raw = str(inc or "").strip()
    match = re.search(r"(\d{5,})", raw)
    return match.group(1) if match else raw


def drive_folder_url(folder_id: str) -> str:
    folder_id = str(folder_id or "").strip()
    if not folder_id:
        return ""
    return f"https://drive.google.com/drive/folders/{folder_id}"


def write_suite_headers(path: Path, metadata: Dict[str, Any], *, case_count: int = 0) -> int:
    from openpyxl import load_workbook

    anchor = load_anchor_map()
    headers = anchor.get("suite_header_cells") or {}
    if not headers:
        return 0
    wb = load_workbook(path)
    written = 0
    suite_title = metadata.get("suite_title") or metadata.get("objetivo") or ""
    section = metadata.get("suite_section") or "CASOS DE PRUEBA"
    if case_count:
        section = f"{section} ({case_count})"
    mapping = {
        "suite_title": suite_title,
        "suite_section": section,
    }
    for field, ref in headers.items():
        value = mapping.get(field)
        if not value or "!" not in ref:
            continue
        sheet, cell = ref.split("!", 1)
        if sheet in wb.sheetnames:
            _set_cell_value(wb[sheet], cell, value)
            written += 1
    wb.save(path)
    wb.close()
    return written


def build_alcance_text(
    obligations: List[Dict[str, Any]],
    *,
    project_id: str = "",
    run_id: str = "",
    summary: str = "",
) -> str:
    """Build scope text for Datos Generales!M29 from interpreted obligations."""
    parts: List[str] = []
    if summary.strip():
        parts.append(summary.strip())
    if project_id:
        parts.append(f"Proyecto APTP: {project_id}")
    if run_id:
        parts.append(f"Ejecución APTP: {run_id}")
    if obligations:
        if parts:
            parts.append("")
        parts.append("Casos incluidos en el set de pruebas:")
        for ob in obligations:
            if not isinstance(ob, dict):
                continue
            title = (ob.get("title") or ob.get("id") or "").strip()
            if title:
                tier = ob.get("automation_tier", "partial")
                parts.append(f"• {title} [{automation_tier_label_es(tier)}]")
    return "\n".join(parts)[:8000]


def _write_hyperlink_cell(ws, cell_ref: str, text: str, url: str) -> bool:
    if not url:
        return False
    cell = _writable_cell(ws, cell_ref)
    cell.value = sanitize_excel_text(text)
    cell.hyperlink = url
    try:
        from openpyxl.styles import Font

        cell.font = Font(color="0563C1", underline="single")
    except Exception:
        pass
    return True


def write_general_data(path: Path, metadata: Dict[str, Any]) -> int:
    from openpyxl import load_workbook

    anchor = load_anchor_map()
    cells_map = anchor.get("general_data_cells", {})
    inc_raw = metadata.get("inc")
    field_values = {
        "inc": _normalize_inc_number(inc_raw) if inc_raw else None,
        "tec": metadata.get("tec"),
        "objetivo": metadata.get("objetivo"),
        "alcance": metadata.get("alcance"),
        "app_platform": metadata.get("app_platform"),
        "start_date": metadata.get("start_date"),
        "cert_date": metadata.get("cert_date"),
    }
    wb = load_workbook(path)
    cleared = 0
    cells_to_clear = anchor.get("boilerplate_cells_to_clear") or []
    cleared += _clear_boilerplate_cells(wb, cells_to_clear)
    written = 0
    for field, ref in cells_map.items():
        if field == "documentation_link":
            continue
        value = field_values.get(field)
        if not value or "!" not in ref:
            continue
        sheet, cell = ref.split("!", 1)
        if sheet in wb.sheetnames:
            _set_cell_value(wb[sheet], cell, value)
            written += 1
    doc_ref = cells_map.get("documentation_link")
    doc_url = metadata.get("documentation_url") or metadata.get("drive_folder_url")
    if doc_ref and doc_url and "!" in doc_ref:
        sheet, cell = doc_ref.split("!", 1)
        if sheet in wb.sheetnames:
            label = metadata.get("documentation_label") or "Documentación técnica APTP (Google Drive)"
            if _write_hyperlink_cell(wb[sheet], cell, label, doc_url):
                written += 1
    wb.save(path)
    wb.close()
    return written + cleared


def _obligation_match_tokens(obligation: Dict[str, Any]) -> Set[str]:
    tokens: Set[str] = set()
    for field in ("id", "title", "external_id"):
        raw = str(obligation.get(field) or "")
        if not raw:
            continue
        tokens.add(raw.casefold())
        for match in re.finditer(r"EMPRESARIAL[-\s]?(\d+)", raw, re.I):
            tokens.add(f"empresarial-{match.group(1)}")
        for match in re.finditer(r"OBL-EMP-(\d+)", raw, re.I):
            tokens.add(f"obl-emp-{match.group(1)}")
            tokens.add(f"empresarial-{match.group(1)}")
    return {t for t in tokens if t}


def _haystack_matches_tokens(haystack: str, tokens: Set[str]) -> bool:
    norm = haystack.casefold().replace(" ", "").replace("_", "-")
    for token in tokens:
        tnorm = token.replace(" ", "").replace("_", "-")
        if tnorm and tnorm in norm:
            return True
    return False


def _find_executed_test(
    obligation: Dict[str, Any],
    execution_results: Dict[str, Any],
    *,
    scenario: Optional[Dict[str, Any]] = None,
) -> Optional[Dict[str, Any]]:
    tokens = _obligation_match_tokens(obligation)
    if scenario:
        for field in ("id", "name", "title"):
            raw = str(scenario.get(field) or "")
            if raw:
                tokens.add(raw.casefold())
                for match in re.finditer(r"EMPRESARIAL[-\s]?(\d+)", raw, re.I):
                    tokens.add(f"empresarial-{match.group(1)}")
    if not tokens:
        return None
    for test in execution_results.get("executed_tests") or []:
        if not isinstance(test, dict):
            continue
        hay = " ".join(
            str(test.get(k) or "") for k in ("id", "title", "file", "name")
        )
        if _haystack_matches_tokens(hay, tokens):
            return test
        if scenario:
            for req in scenario.get("requirement_ids") or []:
                req_slug = str(req).casefold().replace("_", "-")
                if req_slug and req_slug in hay.casefold().replace("_", "-"):
                    return test
    stdout = execution_results.get("stdout") or ""
    for token in tokens:
        emp_token = token.replace("empresarial-", "EMPRESARIAL-")
        if "empresarial" in token and emp_token in stdout:
            line_hit = any(
                emp_token in line and ("✘" in line or "failed" in line.casefold())
                for line in stdout.splitlines()
            )
            if line_hit or execution_results.get("status") in ("failed", "error", "partial"):
                status = "failed" if line_hit or execution_results.get("failed") else "unknown"
                return {"status": status, "title": token, "stdout_match": True}
    return None


def _scenario_for_obligation(
    obligation: Dict[str, Any],
    scenario_by_ob: Dict[str, Dict[str, Any]],
) -> Dict[str, Any]:
    ob_id = str(obligation.get("id") or "")
    if ob_id and ob_id in scenario_by_ob:
        return scenario_by_ob[ob_id]
    tokens = _obligation_match_tokens(obligation)
    for key, sc in scenario_by_ob.items():
        hay = " ".join(
            str(sc.get(k) or "") for k in ("id", "name", "title", "external_id")
        )
        if _haystack_matches_tokens(hay, tokens):
            return sc
        for req in sc.get("requirement_ids") or sc.get("obligation_ids") or []:
            if str(req) in tokens or str(req).casefold() in tokens:
                return sc
    return {}


def _scenario_steps_text(scenario: Dict[str, Any]) -> List[str]:
    lines: List[str] = []
    for step in scenario.get("steps") or []:
        if isinstance(step, dict):
            action = step.get("action") or ""
            target = step.get("target") or ""
            data = step.get("data") or ""
            line = f"{action} {target} {data}".strip()
            if line:
                lines.append(line)
        elif step:
            lines.append(_step_line(step))
    return lines


def _coerce_mapping(value: Any) -> Dict[str, Any]:
    return value if isinstance(value, dict) else {}


def _step_line(step: Any) -> str:
    if isinstance(step, str):
        return step.strip()
    if isinstance(step, dict):
        return f"{step.get('action', '')} {step.get('target', '')}".strip()
    return str(step).strip() if step is not None else ""


def _assertion_line(assertion: Any) -> str:
    if isinstance(assertion, str):
        return assertion.strip()
    if isinstance(assertion, dict):
        return f"{assertion.get('type', '')}: {assertion.get('expected', '')}".strip()
    return str(assertion).strip() if assertion is not None else ""


def _format_test_row_cell(row: Dict[str, str]) -> str:
    title = row.get("cp_code") or row.get("title", row.get("text", ""))
    parts = [title] if title else []
    tier = row.get("automation_tier", "")
    if tier == "manual":
        parts.append("Clasificación: Operación manual (no automatizable)")
        if row.get("manual_rationale"):
            parts.append(f"Motivo: {row['manual_rationale']}")
    elif tier:
        parts.append(f"Clasificación: {automation_tier_label_es(tier)}")
    if row.get("steps"):
        parts.append(f"Pasos:\n{row['steps']}")
    if row.get("expected"):
        parts.append(f"Resultado esperado:\n{row['expected']}")
    result = row.get("result", row.get("status", ""))
    if result:
        parts.append(f"Resultado ejecución: {result}")
    if row.get("evidence"):
        parts.append(f"Evidencia: {row['evidence']}")
    return "\n\n".join(parts)


def scrub_template_sample_content(path: Path) -> int:
    """Remove known E59 template sample fingerprints from all sheets."""
    from openpyxl import load_workbook

    anchor = load_anchor_map()
    fingerprints = anchor.get("template_fingerprints") or []
    if not fingerprints:
        return 0

    wb = load_workbook(path)
    cleared = 0
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows():
            for cell in row:
                val = str(cell.value or "")
                if not val:
                    continue
                if any(fp and fp in val for fp in fingerprints):
                    cell.value = None
                    if getattr(cell, "hyperlink", None):
                        cell.hyperlink = None
                    cleared += 1
    wb.save(path)
    wb.close()
    return cleared


def _column_indices(columns: List[str]) -> List[int]:
    from openpyxl.utils import column_index_from_string

    return [column_index_from_string(col) for col in columns]


def _clear_boilerplate_cells(wb, cells_to_clear: List[str]) -> int:
    cleared = 0
    for ref in cells_to_clear:
        if "!" not in ref:
            continue
        sheet, cell = ref.split("!", 1)
        if sheet not in wb.sheetnames:
            continue
        target = _writable_cell(wb[sheet], cell)
        if target.value not in (None, ""):
            target.value = None
            cleared += 1
        if getattr(target, "hyperlink", None):
            target.hyperlink = None
            cleared += 1
    return cleared


def clear_test_case_rows(
    path: Path,
    *,
    start_row: Optional[int] = None,
    end_row: Optional[int] = None,
    columns: Optional[List[str]] = None,
    boilerplate_cells: Optional[List[str]] = None,
) -> int:
    """Remove example/template rows from Casos de Pruebas before writing run data."""
    from openpyxl import load_workbook

    anchor = load_anchor_map()
    sheet_name = "Casos de Pruebas"
    start = start_row if start_row is not None else anchor.get("test_rows_start", 12) - 2
    end = end_row if end_row is not None else anchor.get("test_rows_end", 500)
    col_letters = columns or anchor.get("test_row_columns") or ["C", "D", "E", "F", "G", "H"]
    col_indices = _column_indices(col_letters)
    cells_to_clear = boilerplate_cells if boilerplate_cells is not None else anchor.get("boilerplate_cells_to_clear") or []

    wb = load_workbook(path)
    cleared = 0
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row_idx in range(start, end + 1):
            for col_idx in col_indices:
                cell = ws.cell(row=row_idx, column=col_idx)
                if cell.value not in (None, ""):
                    cell.value = None
                    cleared += 1
    cleared += _clear_boilerplate_cells(wb, cells_to_clear)
    wb.save(path)
    wb.close()
    return cleared


def _first_compile_message(execution_results: Dict[str, Any]) -> str:
    failures = execution_results.get("failures") or []
    for failure in failures:
        if not isinstance(failure, dict):
            continue
        err = failure.get("error")
        if isinstance(err, dict):
            message = (err.get("message") or "").strip()
        else:
            message = str(err or "").strip()
        if message:
            return message.split("\n", 1)[0][:500]
    stderr = (execution_results.get("stderr") or execution_results.get("stdout") or "").strip()
    if stderr:
        return stderr.split("\n", 1)[0][:500]
    return "Playwright compile or execution failed before tests ran"


def _failure_for_obligation(
    obligation: Dict[str, Any],
    failures: List[Any],
) -> Optional[str]:
    ob_id = obligation.get("id", "")
    title = (obligation.get("title") or "").casefold()
    for failure in failures:
        if not isinstance(failure, dict):
            continue
        fail_title = (failure.get("title") or "").casefold()
        if ob_id and ob_id in (failure.get("title") or failure.get("id") or ""):
            pass
        elif title and title not in fail_title and fail_title not in title:
            continue
        err = failure.get("error")
        if isinstance(err, dict):
            message = (err.get("message") or "").strip()
        else:
            message = str(err or "").strip()
        if message:
            return message.split("\n", 1)[0][:500]
    return None


def _execution_result_label(
    obligation: Dict[str, Any],
    execution_results: Optional[Dict[str, Any]],
    *,
    scenario: Optional[Dict[str, Any]] = None,
) -> Tuple[str, str]:
    """Return (result label, evidence note)."""
    execution_results = _coerce_mapping(execution_results)
    match = _find_executed_test(obligation, execution_results, scenario=scenario)

    tier = obligation.get("automation_tier", "partial")
    has_automation = tier in ("automated", "partial") or bool(scenario)

    # Infrastructure-blocked runs never produced valid functional verdicts, so
    # failures are environment issues rather than product rejections.
    infra_blocked = bool(execution_results.get("infrastructure_blocker")) or (
        str(execution_results.get("release_verdict") or "").lower() == "invalid"
    )

    if match:
        status = (match.get("status") or "unknown").casefold()
        if status in ("passed", "expected"):
            return "Aprobado", ""
        if status == "blocked" or (
            infra_blocked and status in ("failed", "unexpected", "error", "unknown")
        ):
            return "Bloqueado (entorno)", ""
        if status in ("failed", "unexpected", "error", "unknown") and (
            status != "unknown" or execution_results.get("status") in ("failed", "error")
        ):
            err = match.get("error")
            if isinstance(err, dict):
                note = sanitize_excel_text(err.get("message") or "", max_len=500)
            else:
                note = sanitize_excel_text(err or "", max_len=500)
            if not note and match.get("stdout_match"):
                note = "Prueba Playwright fallida (ver evidencia de ejecución)"
            return "Rechazado", note

    if has_automation:
        ob_id = obligation.get("id", "")
        title = (obligation.get("title") or "").casefold()
        exec_by_id = {
            t.get("id", ""): t for t in execution_results.get("executed_tests") or [] if isinstance(t, dict)
        }
        exec_by_title = {
            (t.get("title") or "").casefold(): t
            for t in execution_results.get("executed_tests") or []
            if isinstance(t, dict)
        }
        legacy = exec_by_id.get(ob_id) or exec_by_title.get(title)
        if legacy:
            if legacy.get("status") == "passed":
                return "Aprobado", ""
            return ("Bloqueado (entorno)" if infra_blocked else "Rechazado"), ""

    executed_tests = execution_results.get("executed_tests") or []
    status = (execution_results.get("status") or "").casefold()
    failures = execution_results.get("failures") or []

    if tier == "manual" and not scenario and not match:
        return "Manual", ""

    if infra_blocked and (status in ("failed", "error") or failures or (has_automation and not executed_tests)):
        return "No ejecutado", "entorno: ejecución bloqueada por infraestructura"

    if status in ("failed", "error") and not executed_tests:
        compile_msg = _failure_for_obligation(obligation, failures) or _first_compile_message(execution_results)
        return "Rechazado", f"compilación: {compile_msg}"

    if failures:
        compile_msg = _failure_for_obligation(obligation, failures)
        if compile_msg:
            return "Rechazado", f"compilación: {compile_msg}"

    if has_automation and status in ("failed", "error"):
        return "Rechazado", "ejecución: sin resultado de prueba coincidente"

    return "Pendiente", ""


def build_deliverable_rows(
    obligations: List[Dict[str, Any]],
    *,
    functional_scenarios: Optional[List[Dict[str, Any]]] = None,
    execution_results: Optional[Dict[str, Any]] = None,
    run_id: str = "",
) -> List[Dict[str, str]]:
    """Map interpreted obligations + scenarios to E59 row cells."""
    scenarios = functional_scenarios or []
    scenario_by_ob: Dict[str, Dict[str, Any]] = {}
    for sc in scenarios:
        for req_id in sc.get("requirement_ids") or sc.get("obligation_ids") or []:
            scenario_by_ob[str(req_id)] = sc
        oid = sc.get("id") or sc.get("obligation_id")
        if oid:
            scenario_by_ob[str(oid)] = sc
        ext = sc.get("external_id")
        if ext:
            scenario_by_ob[str(ext)] = sc

    rows: List[Dict[str, str]] = []
    coded_obligations = assign_case_codes([o for o in obligations if isinstance(o, dict)])
    for ob in coded_obligations:
        index_match = coded_obligations.index(ob) + 1
        sc = _scenario_for_obligation(ob, scenario_by_ob)
        steps_parts = list(ob.get("preconditions") or []) + list(ob.get("steps") or [])
        auto_steps = _scenario_steps_text(sc)
        if auto_steps:
            steps_parts = auto_steps
        elif sc.get("steps"):
            steps_parts.extend(_step_line(s) for s in sc.get("steps", [])[:12] if _step_line(s))
        expected_parts = list(ob.get("expected_results") or [])
        if sc.get("assertions"):
            expected_parts.extend(
                line for a in sc.get("assertions", [])[:8] if (line := _assertion_line(a))
            )
        elif sc.get("expected_results"):
            expected_parts.extend(str(e) for e in sc.get("expected_results", []) if e)
        result, compile_evidence = _execution_result_label(ob, execution_results, scenario=sc)
        match = _find_executed_test(ob, _coerce_mapping(execution_results), scenario=sc)
        evidence = compile_evidence
        if match and match.get("evidence_refs"):
            screenshot_evidence = "; ".join(
                ref.get("screenshot_path", "")
                for ref in match["evidence_refs"][:3]
                if isinstance(ref, dict) and ref.get("screenshot_path")
            )
            if screenshot_evidence:
                evidence = screenshot_evidence
        if not evidence and run_id and result in ("Rechazado", "Bloqueado (entorno)"):
            evidence = f"run:{run_id}"
        title = ob.get("title", "")
        cp_code = ob.get("cp_code") or format_case_code(index_match, title)
        rows.append(
            {
                "cp_code": cp_code,
                "title": title,
                "steps": "\n".join(str(s) for s in steps_parts if s)[:4000],
                "expected": "\n".join(str(e) for e in expected_parts if e)[:4000],
                "result": result,
                "evidence": evidence,
                "scenario_type": (sc or {}).get("scenario_type", "happy_path"),
                "obligation_id": ob.get("id", ""),
                "external_id": ob.get("external_id", ""),
                "automation_tier": ob.get("automation_tier", "partial"),
                "manual_rationale": ob.get("automation_rationale", "") if ob.get("automation_tier") == "manual" else "",
            }
        )
    return rows


def append_test_rows(path: Path, rows: List[Dict[str, str]], *, start_row: Optional[int] = None) -> int:
    from openpyxl import load_workbook

    anchor = load_anchor_map()
    sheet_name = "Casos de Pruebas"
    start = start_row or anchor.get("test_rows_start", 12)
    field_cols = anchor.get("test_row_field_columns") or {
        "title": "C",
        "steps": "D",
        "expected": "E",
        "result": "F",
    }
    write_mode = anchor.get("test_row_write_mode") or "merged_cell"
    wb = load_workbook(path)
    if sheet_name not in wb.sheetnames:
        wb.close()
        return 0
    ws = wb[sheet_name]
    appended = 0
    for offset, row in enumerate(rows):
        r = start + offset
        title_col = field_cols.get("title", "C")
        if write_mode == "cp_code":
            _set_cell_value(ws, f"{title_col}{r}", _format_test_row_cell(row))
        elif write_mode == "merged_cell":
            _set_cell_value(ws, f"{title_col}{r}", _format_test_row_cell(row))
        else:
            _set_cell_value(ws, f"{title_col}{r}", row.get("title") or row.get("hierarchy") or "")
            steps_col = field_cols.get("steps", "D")
            if steps_col:
                _set_cell_value(ws, f"{steps_col}{r}", row.get("steps") or "")
            expected_col = field_cols.get("expected", "E")
            if expected_col:
                _set_cell_value(ws, f"{expected_col}{r}", row.get("expected") or "")
            result_col = field_cols.get("result", "F")
            result_val = row.get("result") or ""
            if row.get("evidence") and result_val == "Rechazado":
                result_val = f"{result_val} — {sanitize_excel_text(row['evidence'], max_len=500)}"
            if result_col:
                _set_cell_value(ws, f"{result_col}{r}", result_val)
        appended += 1
    wb.save(path)
    wb.close()
    return appended


def _scan_workbook_fingerprints(wb, fingerprints: List[str]) -> List[str]:
    issues: List[str] = []
    if not fingerprints:
        return issues
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows():
            for cell in row:
                val = str(cell.value or "")
                for fp in fingerprints:
                    if fp and fp in val:
                        issues.append(f"Template fingerprint '{fp}' found at {sheet_name}!{cell.coordinate}")
    return issues


def _validate_cleared_range(
    wb,
    *,
    sheet_name: str,
    start_row: int,
    end_row: int,
    columns: List[str],
    written_start: int,
    written_count: int,
    header_row_end: int = 0,
) -> List[str]:
    issues: List[str] = []
    if sheet_name not in wb.sheetnames:
        return issues
    ws = wb[sheet_name]
    col_indices = _column_indices(columns)
    written_end = written_start + max(written_count, 0) - 1
    for row_idx in range(start_row, end_row + 1):
        if row_idx <= header_row_end:
            continue
        if written_count > 0 and written_start <= row_idx <= written_end:
            continue
        for col_idx in col_indices:
            val = ws.cell(row=row_idx, column=col_idx).value
            if val not in (None, ""):
                from openpyxl.utils import get_column_letter

                coord = f"{get_column_letter(col_idx)}{row_idx}"
                issues.append(f"Stale template content at {sheet_name}!{coord}")
    return issues


def validate_structure(
    path: Path,
    *,
    written_start: Optional[int] = None,
    written_count: int = 0,
) -> Dict[str, Any]:
    from openpyxl import load_workbook

    anchor = load_anchor_map()
    issues: List[str] = []
    wb = load_workbook(path)
    template_wb = load_workbook(_TEMPLATE) if _TEMPLATE.is_file() else None
    expected_sheets = anchor.get("sheets", [])
    for name in expected_sheets:
        if name not in wb.sheetnames:
            issues.append(f"Missing sheet: {name}")
    merge_delta = 0
    expected_merges = anchor.get("merge_counts") or {}
    if template_wb:
        for name in expected_sheets:
            if name in wb.sheetnames and name in template_wb.sheetnames:
                actual = len(list(wb[name].merged_cells.ranges))
                expected = expected_merges.get(name, len(list(template_wb[name].merged_cells.ranges)))
                merge_delta += abs(actual - expected)
    for item in anchor.get("anchors", []):
        sheet, cell, expected = item.get("sheet"), item.get("cell"), item.get("expected_contains", "")
        if sheet in wb.sheetnames and expected:
            val = str(wb[sheet][cell].value or "")
            if expected.lower() not in val.lower():
                issues.append(f"Anchor missing at {sheet}!{cell}")

    fingerprints = anchor.get("template_fingerprints") or []
    issues.extend(_scan_workbook_fingerprints(wb, fingerprints))

    clear_start = anchor.get("test_rows_start", 12) - 2
    clear_end = anchor.get("test_rows_end", 500)
    header_row_end = int(anchor.get("template_header_rows_end", clear_start - 1))
    col_letters = anchor.get("test_row_columns") or ["C", "D", "E", "F", "G", "H"]
    row_start = written_start if written_start is not None else anchor.get("test_rows_start", 12)
    issues.extend(
        _validate_cleared_range(
            wb,
            sheet_name="Casos de Pruebas",
            start_row=clear_start,
            end_row=clear_end,
            columns=col_letters,
            written_start=row_start,
            written_count=written_count,
            header_row_end=int(anchor.get("template_header_rows_end", clear_start - 1)),
        )
    )

    wb.close()
    if template_wb:
        template_wb.close()
    return {"valid": not issues and merge_delta == 0, "issues": issues, "merge_delta": merge_delta}


def build_deliverable(
    dest: Path,
    *,
    metadata: Dict[str, Any],
    obligations: List[Dict[str, Any]],
    execution_results: Optional[Dict[str, Any]] = None,
    functional_scenarios: Optional[List[Dict[str, Any]]] = None,
    run_id: str = "",
) -> Dict[str, Any]:
    clone_template(dest)
    scrub_template_sample_content(dest)
    anchor = load_anchor_map()
    start_row = anchor.get("test_rows_start", 12)
    clear_end = anchor.get("test_rows_end", 500)
    cleared = clear_test_case_rows(dest, start_row=start_row - 2, end_row=clear_end)
    general_written = write_general_data(dest, metadata)
    rows = build_deliverable_rows(
        obligations,
        functional_scenarios=functional_scenarios,
        execution_results=execution_results,
        run_id=run_id,
    )
    suite_written = write_suite_headers(dest, metadata, case_count=len(rows))
    expected_rows = len([o for o in obligations if isinstance(o, dict)])
    appended = append_test_rows(dest, rows, start_row=start_row)
    validation = validate_structure(dest, written_start=start_row, written_count=appended)
    if appended != expected_rows:
        validation = {
            **validation,
            "valid": False,
            "issues": list(validation.get("issues") or [])
            + [f"Expected {expected_rows} test rows, wrote {appended}"],
        }
    return {
        "path": str(dest),
        "general_cells_written": general_written + suite_written,
        "test_rows_appended": appended,
        "obligation_count": expected_rows,
        "template_rows_cleared": cleared,
        "validation": validation,
    }
