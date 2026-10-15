"""Generic Excel extraction — multi-sheet rows without schema assumptions."""
from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List, Optional


def _cell_value(value: Any) -> str:
    if value is None:
        return ""
    return str(value)


def extract_xlsx(path: Path, *, max_rows_per_sheet: int = 5000) -> Dict[str, Any]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    sheets: List[Dict[str, Any]] = []
    text_parts: List[str] = []

    for ws in wb.worksheets:
        rows: List[List[str]] = []
        for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
            if row_idx >= max_rows_per_sheet:
                break
            cells = [_cell_value(c) for c in row]
            if any(cells):
                rows.append(cells)
        max_cols = max((len(r) for r in rows), default=0)
        headers: Optional[List[str]] = None
        if rows:
            headers = rows[0]
        sheet_meta = {
            "name": ws.title,
            "row_count": len(rows),
            "col_count": max_cols,
            "headers": headers,
            "rows": rows,
            "dimensions": f"A1:{_col_letter(max_cols)}{max(len(rows), 1)}",
        }
        sheets.append(sheet_meta)
        preview = "\n".join("\t".join(r) for r in rows[:50])
        text_parts.append(f"Sheet: {ws.title}\n{preview}")

    wb.close()
    return {
        "format": "xlsx",
        "sheet_count": len(sheets),
        "sheets": sheets,
        "text": "\n\n".join(text_parts),
        "metadata": {"source": path.name},
    }


def _col_letter(n: int) -> str:
    if n <= 0:
        return "A"
    result = ""
    while n > 0:
        n, rem = divmod(n - 1, 26)
        result = chr(65 + rem) + result
    return result or "A"
