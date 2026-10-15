"""Unified document extractor for intake files."""
from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Dict

from src.intake.extraction.excel_extractor import extract_xlsx
from src.intake.extraction.pdf_extractor import extract_pdf

SENSITIVE_FIELD_HINTS = re.compile(
    r"(password|passwd|clave|token|secret|otp|pin|credential|api[_-]?key)",
    re.I,
)
STEP_HINT = re.compile(r"(pasos?|steps?|precondici|expected|resultado|caso de prueba)", re.I)


def infer_intake_role(filename: str, extraction: Dict[str, Any]) -> str:
    name = filename.lower()
    if "entregable" in name and name.endswith(".xlsx"):
        return "SUPPLEMENTARY"
    text = (extraction.get("text") or "")[:3000].lower()
    fmt = extraction.get("format", "")
    if fmt == "pdf" and extraction.get("page_count", 0) > 10 and STEP_HINT.search(text):
        return "PRIMARY_SPEC"
    if fmt == "xlsx" and SENSITIVE_FIELD_HINTS.search(text):
        return "ASSET_BUNDLE"
    if fmt in {"xlsx", "csv"} and extraction.get("sheet_count", 0) > 1:
        return "ASSET_BUNDLE"
    if STEP_HINT.search(text):
        return "PRIMARY_SPEC"
    return "SUPPLEMENTARY"


class DocumentExtractor:
    """Extract rich structure from intake files without business semantics."""

    def extract(self, path: Path) -> Dict[str, Any]:
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            return extract_pdf(path)
        if suffix in {".xlsx", ".xls"}:
            return extract_xlsx(path)
        if suffix in {".txt", ".md", ".csv"}:
            text = path.read_text(encoding="utf-8", errors="replace")
            return {
                "format": suffix.lstrip("."),
                "text": text,
                "pages": [{"page_number": 1, "text": text, "char_count": len(text), "tables": []}],
                "sheets": [],
                "tables": [],
                "metadata": {"source": path.name},
            }
        if suffix in {".docx", ".doc"}:
            return self._extract_docx(path)
        text = path.read_text(encoding="utf-8", errors="replace")
        return {
            "format": "unknown",
            "text": text,
            "pages": [],
            "sheets": [],
            "tables": [],
            "metadata": {"source": path.name},
        }

    def _extract_docx(self, path: Path) -> Dict[str, Any]:
        try:
            from docx import Document

            doc = Document(path)
            text = "\n".join(p.text for p in doc.paragraphs)
            tables = []
            for table in doc.tables:
                rows = [[cell.text for cell in row.cells] for row in table.rows]
                tables.append({"rows": rows, "row_count": len(rows)})
            return {
                "format": "docx",
                "text": text,
                "pages": [{"page_number": 1, "text": text, "char_count": len(text), "tables": []}],
                "sheets": [],
                "tables": tables,
                "metadata": {"source": path.name},
            }
        except Exception:
            text = path.read_text(encoding="utf-8", errors="replace")
            return {"format": "docx", "text": text, "pages": [], "sheets": [], "tables": [], "metadata": {}}


    def extract_bytes(self, filename: str, content: bytes) -> Dict[str, Any]:
        import tempfile

        suffix = Path(filename).suffix.lower() or ".txt"
        with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
            tmp.write(content)
            tmp_path = Path(tmp.name)
        try:
            result = self.extract(tmp_path)
            result["filename"] = filename
            return result
        finally:
            tmp_path.unlink(missing_ok=True)


def extract_document(path: Path) -> Dict[str, Any]:
    return DocumentExtractor().extract(path)
