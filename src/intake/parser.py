"""Document parsing for PDF, DOCX, XLSX, images, and SVG."""
import io
import re
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from src.intake.extraction.document_extractor import DocumentExtractor
from src.intake.file_types import IMAGE_EXTENSIONS

from src.intake.schemas.canonical import (
    CanonicalDocument, CanonicalSection, DocumentMetadata, DocumentTypeEnum,
)

_extractor = DocumentExtractor()

# Soft intake roles (non-blocking classification)
SOFT_ROLES = ("PRIMARY_SPEC", "ASSET_BUNDLE", "SUPPLEMENTARY", "FIXTURE_FILE")


def infer_document_type(filename: str, text: str) -> DocumentTypeEnum:
    name = filename.upper()
    text_head = text[:2000].upper()
    if "BRD" in name or "BUSINESS REQUIREMENTS" in text_head:
        return DocumentTypeEnum.BRD
    if "PRD" in name or "PRODUCT REQUIREMENTS" in text_head:
        return DocumentTypeEnum.PRD
    if "USE CASE" in text_head or "UC-" in name or "USECASE" in name:
        return DocumentTypeEnum.UC
    if "ACCEPTANCE" in text_head or "AC-" in name:
        return DocumentTypeEnum.AC
    return DocumentTypeEnum.OTHER


def extract_sections(text: str) -> List[CanonicalSection]:
    sections: List[CanonicalSection] = []
    current_title = "Introduction"
    current_lines: List[str] = []
    for line in text.splitlines():
        if re.match(r"^#{1,3}\s+", line) or re.match(r"^[A-Z][A-Z0-9\s\-]{3,}$", line.strip()):
            if current_lines:
                sections.append(CanonicalSection(title=current_title, content="\n".join(current_lines).strip()))
            current_title = line.strip("# ").strip()
            current_lines = []
        else:
            current_lines.append(line)
    if current_lines:
        sections.append(CanonicalSection(title=current_title, content="\n".join(current_lines).strip()))
    return sections or [CanonicalSection(title="Body", content=text)]


class DocumentParser:
    def parse_file(self, path: Path, doc_id: Optional[str] = None) -> CanonicalDocument:
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            text, tables = self._parse_pdf(path)
        elif suffix in {".docx", ".doc"}:
            text, tables = self._parse_docx(path)
        elif suffix in {".xlsx", ".xls"}:
            text, tables = self._parse_xlsx(path)
        elif suffix in {".txt", ".md"}:
            text = path.read_text(encoding="utf-8", errors="replace")
            tables = []
        elif suffix == ".svg":
            text, tables = self._parse_svg(path.read_bytes())
        elif suffix in IMAGE_EXTENSIONS:
            text, tables = self._parse_image(path.name, path.read_bytes())
        else:
            text = path.read_text(encoding="utf-8", errors="replace")
            tables = []

        doc_type = infer_document_type(path.name, text)
        metadata = DocumentMetadata(
            project=path.stem.split("_")[0] if "_" in path.stem else None,
            document_type=doc_type,
        )
        extra = _build_extra_from_path(path, text, tables)
        return CanonicalDocument(
            id=doc_id or path.stem,
            filename=path.name,
            document_type=doc_type,
            metadata=metadata,
            sections=extract_sections(text),
            raw_text=text,
            tables=tables,
            extra=extra,
        )

    def parse_bytes(self, filename: str, content: bytes, doc_id: Optional[str] = None) -> CanonicalDocument:
        import tempfile

        suffix = Path(filename).suffix.lower() or ".txt"
        if suffix == ".svg":
            text, tables = self._parse_svg(content)
            doc_type = infer_document_type(filename, text)
            return CanonicalDocument(
                id=doc_id or Path(filename).stem,
                filename=filename,
                document_type=doc_type,
                metadata=DocumentMetadata(document_type=doc_type),
                sections=extract_sections(text),
                raw_text=text,
                tables=tables,
            )
        if suffix in IMAGE_EXTENSIONS:
            text, tables = self._parse_image(filename, content)
            doc_type = infer_document_type(filename, text)
            return CanonicalDocument(
                id=doc_id or Path(filename).stem,
                filename=filename,
                document_type=doc_type,
                metadata=DocumentMetadata(document_type=doc_type),
                sections=extract_sections(text),
                raw_text=text,
                tables=tables,
            )

        with tempfile.NamedTemporaryFile(suffix=suffix or ".txt", delete=False) as tmp:
            tmp.write(content)
            tmp_path = Path(tmp.name)
        try:
            doc = self.parse_file(tmp_path, doc_id=doc_id)
            doc.filename = filename
            if not doc.extra:
                doc.extra = _build_extra_from_path(tmp_path, doc.raw_text, doc.tables)
            return doc
        finally:
            tmp_path.unlink(missing_ok=True)

    def _parse_pdf(self, path: Path) -> Tuple[str, List]:
        extracted = _extractor.extract(path)
        text = extracted.get("text", "")
        tables = [t.get("rows", t) if isinstance(t, dict) else t for t in extracted.get("tables", [])]
        return text, tables

    def _parse_docx(self, path: Path) -> Tuple[str, List]:
        try:
            from docx import Document
            doc = Document(path)
            text = "\n".join(p.text for p in doc.paragraphs)
            tables = []
            for table in doc.tables:
                rows = [[cell.text for cell in row.cells] for row in table.rows]
                tables.append(rows)
            return text, tables
        except Exception:
            return path.read_text(encoding="utf-8", errors="replace"), []

    def _parse_xlsx(self, path: Path) -> Tuple[str, List]:
        try:
            extracted = _extractor.extract(path)
            text = extracted.get("text", "")
            tables = [s.get("rows", []) for s in extracted.get("sheets", [])]
            return text, tables
        except Exception:
            return "", []

    def _parse_svg(self, content: bytes) -> Tuple[str, List]:
        text = content.decode("utf-8", errors="replace")
        text = re.sub(r"<script[^>]*>.*?</script>", "", text, flags=re.I | re.S)
        visible = re.sub(r"<[^>]+>", " ", text)
        visible = re.sub(r"\s+", " ", visible).strip()
        header = f"SVG asset ({len(content)} bytes).\n"
        body = visible[:100_000] if visible else "(no extractable text in SVG)"
        return header + body, []

    def _parse_image(self, filename: str, content: bytes) -> Tuple[str, List]:
        lines = [
            f"Image asset: {filename}",
            f"Size: {len(content)} bytes",
        ]
        try:
            from PIL import Image

            img = Image.open(io.BytesIO(content))
            lines.append(f"Format: {img.format}, Dimensions: {img.size[0]}x{img.size[1]}")
            if img.mode:
                lines.append(f"Color mode: {img.mode}")
        except Exception:
            lines.append("Could not decode image; kept as binary UX reference.")
        lines.append("Use this asset as UX/wireframe/visual requirements context.")
        return "\n".join(lines), []


_parser = DocumentParser()


def detect_doc_type(filename: str, content: str = "") -> DocumentTypeEnum:
    return infer_document_type(filename, content)


def parse_text_file(path: Path, project_id: str, document_id: str) -> CanonicalDocument:
    doc = _parser.parse_file(Path(path), doc_id=document_id)
    doc.metadata.project = project_id
    return doc


def parse_bytes(
    data: bytes, filename: str, project_id: str, document_id: str
) -> CanonicalDocument:
    doc = _parser.parse_bytes(filename, data, doc_id=document_id)
    doc.metadata.project = project_id
    return doc


def infer_soft_role(filename: str, text: str = "", sheet_count: int = 0) -> str:
    """Heuristic soft classification — non-blocking."""
    name = filename.lower()
    if "entregable" in name or "set de pruebas" in name:
        return "SUPPLEMENTARY"
    if sheet_count > 1 or any(tok in name for tok in ("data", "usuarios", "credencial", "credential", "inscripcion", "test-credentials")):
        return "ASSET_BUNDLE"
    if len(text) > 15000 or any(tok in text.lower() for tok in ("pasos:", "precondicion", "test case", "caso de prueba")):
        return "PRIMARY_SPEC"
    if name.endswith((".txt", ".csv", ".dat")):
        return "FIXTURE_FILE"
    return "SUPPLEMENTARY"


def _build_extra_from_path(path: Path, text: str, tables: List) -> Dict[str, Any]:
    suffix = path.suffix.lower()
    extra: Dict[str, Any] = {"format": suffix.lstrip(".")}
    if suffix == ".pdf":
        extracted = _extractor.extract(path)
        extra["pages"] = extracted.get("pages", [])
        extra["tables"] = extracted.get("tables", [])
        extra["metadata"] = {
            "page_count": extracted.get("page_count", len(extra["pages"])),
            "has_steps": any(
                tok in text.lower()[:8000]
                for tok in ("pasos:", "precondicion", "test case", "caso de prueba")
            ),
            "has_urls": "http://" in text.lower() or "https://" in text.lower(),
        }
    elif suffix in {".xlsx", ".xls"}:
        extracted = _extractor.extract(path)
        sheets = extracted.get("sheets", [])
        extra["sheets"] = sheets
        extra["sheet_metadata"] = {
            s.get("name", f"sheet_{i}"): {
                "dimensions": {"rows": s.get("row_count", 0), "cols": s.get("col_count", 0)},
                "headers": (s.get("headers") or [])[:30],
            }
            for i, s in enumerate(sheets)
        }
        extra["metadata"] = {
            "sheet_count": extracted.get("sheet_count", len(sheets)),
            "has_password_hints": any(
                tok in text.lower()
                for tok in ("clave", "password", "token", "secret")
            ),
        }
        extra["tables"] = tables
    else:
        extra["tables"] = tables
        extra["metadata"] = {}
    extra["intake_role"] = infer_soft_role(
        path.name, text, sheet_count=extra.get("metadata", {}).get("sheet_count", 0)
    )
    return extra
