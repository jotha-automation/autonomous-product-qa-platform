"""Tests for Excel deliverable writer."""
from pathlib import Path

import pytest

from src.intake.test_naming import format_case_code
from src.reporting.excel_deliverable_report import (
    build_alcance_text,
    build_deliverable,
    drive_folder_url,
    load_anchor_map,
    validate_structure,
)

TEMPLATE = Path("src/intake/templates/deliverables/set-de-pruebas-template.xlsx")


def _test_row_start() -> int:
    return int(load_anchor_map().get("test_rows_start", 12))


def _case_row(index: int = 0) -> int:
    return _test_row_start() + index


def _workbook_has_fingerprint(path: Path, fingerprint: str) -> bool:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True)
    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row in ws.iter_rows(min_row=10, max_row=500):
                for cell in row:
                    if fingerprint in str(cell.value or ""):
                        return True
    finally:
        wb.close()
    return False


def _cleared_rows_empty(path: Path, *, start: int = 10, end: int = 500, written_start: int = 12, written_count: int) -> bool:
    from openpyxl import load_workbook
    from openpyxl.utils import column_index_from_string

    anchor = load_anchor_map()
    columns = anchor.get("test_row_columns") or ["C", "D", "E", "F", "G", "H"]
    col_indices = [column_index_from_string(c) for c in columns]
    written_end = written_start + max(written_count, 0) - 1
    header_row_end = int(anchor.get("template_header_rows_end", written_start - 1))

    wb = load_workbook(path, read_only=True)
    try:
        ws = wb["Casos de Pruebas"]
        for row_idx in range(start, end + 1):
            if row_idx <= header_row_end:
                continue
            if written_count > 0 and written_start <= row_idx <= written_end:
                continue
            for col_idx in col_indices:
                if ws.cell(row=row_idx, column=col_idx).value not in (None, ""):
                    return False
    finally:
        wb.close()
    return True


@pytest.mark.skipif(not TEMPLATE.is_file(), reason="E59 template not present")
def test_build_deliverable_clears_example_rows(tmp_path):
    from openpyxl import load_workbook

    dest = tmp_path / "out.xlsx"
    obligations = [
        {"id": "OBL-001", "title": "Rechazo 1 día hábil", "automation_tier": "automated", "steps": ["Paso 1"]},
    ]
    build_deliverable(
        dest,
        metadata={"inc": "INC1353166", "objetivo": "Inscripciones ACH", "alcance": "Front inscripciones"},
        obligations=obligations,
        execution_results={"executed_tests": [{"id": "OBL-001", "status": "failed"}]},
        run_id="run-test",
    )
    wb = load_workbook(dest, read_only=True)
    ws = wb["Casos de Pruebas"]
    row = _case_row(0)
    cell = str(ws[f"C{row}"].value or "")
    assert "INC1346999" not in cell
    assert cell.startswith("CP_0001_")
    assert "RECHAZO" in cell
    assert "Resultado ejecución: Rechazado" in cell
    wb.close()
    assert not _workbook_has_fingerprint(dest, "INC1346999")
    assert _cleared_rows_empty(dest, written_start=_test_row_start(), written_count=len(obligations))


@pytest.mark.skipif(not TEMPLATE.is_file(), reason="E59 template not present")
def test_build_deliverable_no_template_fingerprints(tmp_path):
    obligations = [
        {"id": "OBL-001", "title": "Validar login", "automation_tier": "automated"},
        {"id": "OBL-002", "title": "Validar flujo manual", "automation_tier": "manual"},
    ]
    dest = tmp_path / "out.xlsx"
    result = build_deliverable(
        dest,
        metadata={"inc": "INC999", "objetivo": "Test objective", "alcance": "Scope A"},
        obligations=obligations,
        execution_results={"executed_tests": [{"id": "OBL-001", "status": "passed"}]},
    )
    assert result["validation"]["valid"] is True
    assert not _workbook_has_fingerprint(dest, "INC1346999")
    assert _cleared_rows_empty(dest, written_start=_test_row_start(), written_count=len(obligations))


@pytest.mark.skipif(not TEMPLATE.is_file(), reason="E59 template not present")
def test_build_deliverable_marks_compile_fail_on_empty_execution(tmp_path):
    dest = tmp_path / "out.xlsx"
    obligations = [
        {"id": "OBL-001", "title": "Crear inscripción", "automation_tier": "automated"},
        {"id": "OBL-002", "title": "Revisión manual", "automation_tier": "manual"},
    ]
    build_deliverable(
        dest,
        metadata={"inc": "INC999", "objetivo": "Test", "alcance": "Scope"},
        obligations=obligations,
        execution_results={
            "status": "failed",
            "executed_tests": [],
            "failures": [
                {
                    "title": "Crear inscripción",
                    "error": {"message": "SyntaxError: Identifier 'captureStepEvidence' has already been declared"},
                }
            ],
        },
        run_id="run-compile-fail",
    )
    from openpyxl import load_workbook

    wb = load_workbook(dest, read_only=True)
    ws = wb["Casos de Pruebas"]
    row = _case_row(0)
    automated_cell = str(ws[f"C{row}"].value or "")
    manual_cell = str(ws[f"C{row + 1}"].value or "")
    wb.close()
    assert automated_cell.startswith("CP_0001_")
    assert "Resultado ejecución: Rechazado" in automated_cell
    assert "compilación:" in automated_cell
    assert "captureStepEvidence" in automated_cell
    assert "Resultado ejecución: Manual" in manual_cell
    assert manual_cell.startswith("CP_0002_")
    assert "Pendiente" not in automated_cell
    assert "CREAR_INSCRIPCIÓN" in automated_cell or "CREAR_INSCRIPCION" in automated_cell


@pytest.mark.skipif(not TEMPLATE.is_file(), reason="E59 template not present")
def test_build_deliverable_structure(tmp_path):
    dest = tmp_path / "out.xlsx"
    result = build_deliverable(
        dest,
        metadata={"inc": "INC999", "objetivo": "Test objective", "alcance": "Scope A"},
        obligations=[
            {"id": "OBL-001", "title": "Validar login", "automation_tier": "automated"},
            {"id": "OBL-002", "title": "Validar flujo manual", "automation_tier": "manual"},
        ],
        execution_results={"executed_tests": [{"id": "OBL-001", "status": "passed"}]},
    )
    assert dest.is_file()
    assert result["test_rows_appended"] >= 2
    validation = validate_structure(dest, written_start=_test_row_start(), written_count=2)
    assert validation["merge_delta"] == 0
    assert validation["valid"] is True


@pytest.mark.skipif(not TEMPLATE.is_file(), reason="E59 template not present")
def test_build_deliverable_handles_string_scenario_steps(tmp_path):
    dest = tmp_path / "out.xlsx"
    result = build_deliverable(
        dest,
        metadata={"inc": "INC999", "objetivo": "Test", "alcance": "Scope"},
        obligations=[
            {"id": "OBL-001", "title": "Login flow", "automation_tier": "automated"},
        ],
        functional_scenarios=[
            {
                "id": "OBL-001",
                "requirement_ids": ["OBL-001"],
                "steps": ["Given user on login page", "When user submits credentials"],
                "assertions": ["User reaches dashboard"],
            }
        ],
        execution_results={"executed_tests": [{"id": "OBL-001", "status": "passed"}]},
    )
    assert result["test_rows_appended"] == 1
    from openpyxl import load_workbook

    wb = load_workbook(dest, read_only=True)
    cell = str(wb["Casos de Pruebas"][f"C{_case_row(0)}"].value or "")
    assert "login page" in cell
    wb.close()


@pytest.mark.skipif(not TEMPLATE.is_file(), reason="E59 template not present")
def test_manual_obligation_reflects_playwright_failure_via_scenario(tmp_path):
    dest = tmp_path / "out.xlsx"
    build_deliverable(
        dest,
        metadata={"inc": "INC1353166", "objetivo": "Bank portal", "alcance": "E2E"},
        obligations=[
            {
                "id": "OBL-EMP-13327",
                "title": "EMPRESARIAL-13327: Parámetro = 1 día Hábil",
                "automation_tier": "manual",
                "steps": ["Paso manual del intake"],
                "expected_results": ["Resultado esperado intake"],
            },
        ],
        functional_scenarios=[
            {
                "id": "EMPRESARIAL-13327-P0",
                "name": "Configure processing time parameter to 1 business day",
                "steps": [
                    {"action": "login", "target": "LoginPage", "data": "login_user"},
                    {"action": "testid_fill", "target": "processing-time-param", "data": "1"},
                ],
            }
        ],
        execution_results={
            "status": "failed",
            "passed": 0,
            "failed": 1,
            "executed_tests": [
                {
                    "id": "TC-001",
                    "title": "configure-processing-time-parameter-to-1-business-day-req-001.spec",
                    "status": "unknown",
                }
            ],
            "stdout": "✘ EMPRESARIAL-13327-P0: Configure processing time",
        },
        run_id="run-emp",
    )
    from openpyxl import load_workbook

    wb = load_workbook(dest, read_only=True)
    ws = wb["Casos de Pruebas"]
    cell = str(ws[f"C{_case_row(0)}"].value or "")
    wb.close()
    assert cell.startswith("CP_0001_")
    assert "EMPRESARIAL-13327" in cell or "PARÁMETRO" in cell or "PARAMETRO" in cell
    assert "login" in cell.casefold()
    assert "Resultado ejecución: Rechazado" in cell


@pytest.mark.skipif(not TEMPLATE.is_file(), reason="E59 template not present")
def test_build_deliverable_writes_all_obligation_rows(tmp_path):
    obligations = [
        {"id": f"OBL-{i:03d}", "title": f"Caso de prueba {i}", "automation_tier": "partial", "steps": [f"Paso {i}"]}
        for i in range(1, 33)
    ]
    dest = tmp_path / "out.xlsx"
    result = build_deliverable(
        dest,
        metadata={"inc": "INC999", "objetivo": "Objetivo", "alcance": "Alcance"},
        obligations=obligations,
    )
    assert result["test_rows_appended"] == 32
    assert result["obligation_count"] == 32
    assert result["validation"]["valid"] is True
    from openpyxl import load_workbook

    wb = load_workbook(dest, read_only=True)
    ws = wb["Casos de Pruebas"]
    start = _test_row_start()
    for i in range(32):
        assert ws.cell(row=start + i, column=3).value
    assert ws.cell(row=start + 32, column=3).value in (None, "")
    wb.close()


@pytest.mark.skipif(not TEMPLATE.is_file(), reason="E59 template not present")
def test_write_general_data_clears_departments_and_sets_drive_link(tmp_path):
    from openpyxl import load_workbook

    from src.reporting.excel_deliverable_report import clone_template, write_general_data

    dest = tmp_path / "out.xlsx"
    clone_template(dest)
    folder_id = "1AbCdEfGhIjKlMnOpQrStUvWxYz"
    write_general_data(
        dest,
        {
            "inc": "INC1353166",
            "tec": "TEC-999",
            "objetivo": "Objetivo real del run",
            "alcance": "Alcance detallado del proyecto",
            "documentation_url": drive_folder_url(folder_id),
            "documentation_label": "Documentación técnica APTP (Google Drive)",
        },
    )
    wb = load_workbook(dest)
    dg = wb["Datos Generales"]
    assert dg["C32"].value in (None, "")
    assert dg["G32"].value in (None, "")
    assert dg["C48"].value in (None, "")
    assert dg["G49"].value in (None, "")
    assert "Walter Angulo" not in str(dg["C32"].value or "")
    assert "Alcance detallado" in str(dg["M29"].value or "")
    m49 = dg["M49"]
    assert m49.hyperlink is not None
    assert folder_id in m49.hyperlink.target
    wb.close()


def test_build_alcance_text_lists_obligations():
    obligations = [
        {"id": "OBL-001", "title": "EMPRESARIAL-13327: Caso A", "automation_tier": "manual"},
        {"id": "OBL-002", "title": "EMPRESARIAL-13326: Caso B", "automation_tier": "automated"},
    ]
    text = build_alcance_text(obligations, project_id="PRUEBA", run_id="run-abc")
    assert "Casos incluidos en el set de pruebas" in text
    assert "EMPRESARIAL-13327" in text
    assert "[Operación manual]" in text
    assert "[Automatizable]" in text
    assert "PRUEBA" in text
    assert "run-abc" in text


@pytest.mark.skipif(not TEMPLATE.is_file(), reason="E59 template not present")
def test_write_general_data_uses_template_anchor_cells(tmp_path):
    from openpyxl import load_workbook

    from src.reporting.excel_deliverable_report import clone_template, write_general_data

    dest = tmp_path / "out.xlsx"
    clone_template(dest)
    write_general_data(
        dest,
        {
            "inc": "INC1353166",
            "tec": "TEC-999",
            "objetivo": "Objetivo real del run",
            "alcance": "Alcance real del run",
        },
    )
    wb = load_workbook(dest, read_only=True)
    dg = wb["Datos Generales"]
    assert str(dg["I15"].value) == "1353166"
    assert dg["I16"].value in (None, "")
    assert "Objetivo real del run" in str(dg["M13"].value or "")
    assert "Alcance real del run" in str(dg["M29"].value or "")
    assert "INC1346999" not in str(dg["M13"].value or "")
    wb.close()


@pytest.mark.skipif(not TEMPLATE.is_file(), reason="E59 template not present")
def test_validate_template_fingerprint():
    validation = validate_structure(TEMPLATE)
    assert validation["valid"] is False
    assert any(
        fp in issue
        for issue in validation["issues"]
        for fp in ("CP_0001_CONSULTA SALDOS", "XX/XX/XXXX", "TITULO SUITE PRUEBAS")
    )


@pytest.mark.skipif(not TEMPLATE.is_file(), reason="E59 template not present")
def test_build_deliverable_strips_ansi_from_playwright_evidence(tmp_path):
    dest = tmp_path / "ansi-evidence.xlsx"
    ansi_error = (
        "Error: \x1b[2mexpect(\x1b[22m\x1b[31mlocator\x1b[39m\x1b[2m).\x1b[22m"
        "toBeVisible\x1b[2m(\x1b[22m\x1b[2m)\x1b[22m failed"
    )
    result = build_deliverable(
        dest,
        metadata={"inc": "INC1353166", "objetivo": "Bank portal", "alcance": "E2E"},
        obligations=[
            {
                "id": "OBL-EMP-13330",
                "title": "EMPRESARIAL-13330: Crear Inscripción - Digitar Inscripciones",
                "automation_tier": "automated",
                "steps": ["Crear inscripción"],
                "expected_results": ["Proceso creado"],
            },
        ],
        functional_scenarios=[
            {
                "id": "OBL-EMP-13330",
                "title": "EMPRESARIAL-13330: Crear Inscripción - Digitar Inscripciones",
                "steps": [{"action": "testid_click", "target": "crear-digitar"}],
            }
        ],
        execution_results={
            "status": "failed",
            "failed": 1,
            "executed_tests": [
                {
                    "id": "OBL-EMP-13330",
                    "title": "EMPRESARIAL-13330: Crear Inscripción - Digitar Inscripciones",
                    "status": "failed",
                    "error": {"message": ansi_error},
                }
            ],
        },
        run_id="run-ansi",
    )
    assert result["test_rows_appended"] == 1
    from openpyxl import load_workbook

    wb = load_workbook(dest, read_only=True)
    cell = str(wb["Casos de Pruebas"][f"C{_case_row(0)}"].value or "")
    assert "\x1b" not in cell
    assert "expect(locator).toBeVisible() failed" in cell
    wb.close()


def test_format_case_code_uses_four_digit_index():
    code = format_case_code(1, "EMPRESARIAL-13327: Parámetro = 1 día Hábil")
    assert code.startswith("CP_0001_")
    assert "PARÁMETRO" in code or "PARAMETRO" in code or "DÍA" in code or "DIA" in code

